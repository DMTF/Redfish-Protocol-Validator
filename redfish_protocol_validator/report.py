# Copyright Notice:
# Copyright 2020-2022 DMTF. All rights reserved.
# License: BSD 3-Clause License. For full text see link:
# https://github.com/DMTF/Redfish-Protocol-Validator/blob/main/LICENSE.md

import html as html_mod
import json
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from redfish_service_validator.html_template import build_html_report

from redfish_protocol_validator.constants import Result
from redfish_protocol_validator.system_under_test import SystemUnderTest

sections = [
    ('PROTO_', 'Protocol Details'),
    ('REQ_', 'Service Requests'),
    ('RESP_', 'Service Responses'),
    ('SERV_', 'Service Details'),
    ('SEC_', 'Security Details'),
]

_RPV_EXTRA_CSS = r"""

    /* ── Section ── */
    .section { margin-bottom: 1.5rem; }
    .section-header {
      background: #f7f9fc;
      color: #0d1b2a;
      padding: 10px 16px;
      border-radius: 8px 8px 0 0;
      font-weight: 600;
      font-size: 13px;
      letter-spacing: .01em;
      display: flex;
      align-items: center;
      gap: .55rem;
      cursor: pointer;
      user-select: none;
      border-bottom: 1px solid #dde3ec;
    }
    .section-arrow { font-size: .7rem; color: #6c757d; transition: transform .2s; }
    .section-header.collapsed .section-arrow { transform: rotate(-90deg); }
    .section-body {
      border: 1px solid #dde3ec;
      border-top: none;
      border-radius: 0 0 8px 8px;
      overflow: hidden;
      box-shadow: 0 1px 3px rgba(0,0,0,.05);
      margin-bottom: 8px;
    }

    /* ── Test block ── */
    .test-block { border-bottom: 1px solid #dde3ec; }
    .test-block:last-child { border-bottom: none; }
    .test-heading {
      background: #fff;
      padding: 10px 16px;
      cursor: pointer;
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      gap: 1rem;
    }
    .test-heading:hover { background: #f8fafc; }
    .test-name   { font-weight: 700; font-size: 13px; color: #0d1b2a; font-family: "Cascadia Code", "Consolas", monospace; }
    .test-desc   { font-size: 11px; color: #6c757d; margin-top: .12rem; font-style: italic; }
    .test-toggle { flex-shrink: 0; font-size: .65rem; color: #6c757d; margin-top: .3rem; transition: transform .2s; }
    .test-body        { display: block; }
    .test-body.hidden { display: none; }

    /* ── Results table ── */
    .test-results-table { width: 100%; border-collapse: collapse; font-size: 12px; }
    .test-results-table thead tr { background: #f0f4f8; }
    .test-results-table th {
      padding: 7px 12px;
      text-align: left;
      font-weight: 700;
      font-size: 11px;
      color: #2c3e50;
      border-bottom: 2px solid #dde3ec;
    }
    .test-results-table td {
      padding: 5px 12px;
      border-bottom: 1px solid #f0f2f5;
      vertical-align: top;
      color: #333;
      background: #fff;
    }
    .test-results-table tr:last-child td { border-bottom: none; }
    .test-results-table tr:hover td { background: #f8fafc; }
    .col-method { width:  8%; white-space: nowrap; }
    .col-status { width:  8%; white-space: nowrap; }
    .col-uri    { width: 32%; font-family: "Cascadia Code", "Consolas", monospace; word-break: break-all; }
    .col-result { width:  9%; white-space: nowrap; }
    .col-msg    { width: 43%; }

    /* ── Badges ── */
    .badge {
      display: inline-block;
      padding: .15rem .55rem;
      border-radius: 9999px;
      font-size: .69rem;
      font-weight: 700;
      letter-spacing: .04em;
      text-transform: uppercase;
    }

    /* ── Toolbar ── */
    .toolbar {
      display: flex;
      align-items: center;
      gap: .6rem;
      margin-bottom: 1rem;
    }
    .toolbar-btn {
      background: #0d6efd;
      color: #fff;
      border: none;
      border-radius: 5px;
      padding: 5px 13px;
      font-size: 11px;
      font-weight: 600;
      cursor: pointer;
      letter-spacing: .02em;
      transition: background .15s;
      box-shadow: 0 2px 5px rgba(13,110,253,.3);
    }
    .toolbar-btn:hover { background: #0b5ed7; }

    /* ── Inline result summary chips ── */
    .result-chips { display: flex; gap: .4rem; margin-top: .45rem; flex-wrap: wrap; }
    .rchip {
      display: inline-flex; align-items: center; gap: .3rem;
      padding: .22rem .75rem; border-radius: 9999px;
      font-size: .72rem; font-weight: 700; white-space: nowrap; cursor: default;
    }
    .rchip-pass { background: #c6f6d5; color: #276749; border: 1px solid #38a169; }
    .rchip-warn { background: #fefcbf; color: #744210; border: 1px solid #d69e2e; }
    .rchip-fail { background: #fed7d7; color: #742a2a; border: 1px solid #e53e3e; }
    .rchip-skip { background: #e2e8f0; color: #4a5568; border: 1px solid #a0aec0; }

    /* ── Section count badges ── */
    .section-counts { display: flex; gap: .35rem; margin-left: auto; flex-shrink: 0; align-items: center; }
    .scnt { padding: 2px 8px; border-radius: 20px; font-size: 11px; font-weight: 700; letter-spacing: .01em; }
    .scnt-pass { background: #d4edda; color: #145a32; }
    .scnt-warn { background: #fff3cd; color: #7d6008; }
    .scnt-fail { background: #f8d7da; color: #7b241c; }
    .scnt-skip { background: #f5f7fa; color: #7f8c8d; }

    /* ── Test block status left border ── */
    .test-block.tb-fail { border-left: 4px solid #e53e3e; }
    .test-block.tb-warn { border-left: 4px solid #d69e2e; }
    .test-block.tb-pass { border-left: 4px solid #38a169; }
"""

_RPV_EXTRA_JS = r"""

  /* Section collapse/expand */
  document.querySelectorAll('.section-header').forEach(function(hdr) {
    hdr.addEventListener('click', function() {
      this.classList.toggle('collapsed');
      var body = this.nextElementSibling;
      body.style.display = (body.style.display === 'none') ? '' : 'none';
    });
  });

  /* Test row collapse/expand */
  document.querySelectorAll('.test-heading').forEach(function(hdr) {
    hdr.addEventListener('click', function() {
      var body = this.nextElementSibling;
      var arrow = this.querySelector('.test-toggle');
      if (body && body.classList.contains('test-body')) {
        var hidden = body.classList.toggle('hidden');
        if (arrow) arrow.style.transform = hidden ? 'rotate(-90deg)' : '';
      }
    });
  });

  /* Expand all */
  function expandAll() {
    document.querySelectorAll('.section-header').forEach(function(hdr) {
      hdr.classList.remove('collapsed');
      hdr.nextElementSibling.style.display = '';
    });
    document.querySelectorAll('.test-body').forEach(function(b) {
      b.classList.remove('hidden');
    });
    document.querySelectorAll('.test-toggle').forEach(function(a) {
      a.style.transform = '';
    });
  }

  /* Collapse all */
  function collapseAll() {
    document.querySelectorAll('.section-header').forEach(function(hdr) {
      hdr.classList.add('collapsed');
      hdr.nextElementSibling.style.display = 'none';
    });
    document.querySelectorAll('.test-body').forEach(function(b) {
      b.classList.add('hidden');
    });
    document.querySelectorAll('.test-toggle').forEach(function(a) {
      a.style.transform = 'rotate(-90deg)';
    });
  }

  /* Assertion filter */
  (function() {
    var filter = document.getElementById('uriFilter');
    var clearBtn = document.getElementById('filterClear');
    var countEl = document.getElementById('filterCount');
    if (!filter) return;

    function updateCount() {
      var allTests = document.querySelectorAll('.test-block').length;
      var visibleTests = 0;
      document.querySelectorAll('.test-block').forEach(function(tb) {
        if (tb.style.display !== 'none') visibleTests += 1;
      });
      if (countEl) countEl.innerHTML = '<b>' + visibleTests + '</b> / ' + allTests;
      if (clearBtn) clearBtn.style.display = filter.value.trim() ? 'block' : 'none';
    }

    filter.addEventListener('input', function() {
      var q = (this.value || '').toLowerCase().trim();
      var sections = document.querySelectorAll('.section');

      sections.forEach(function(section) {
        var sectionHdr = section.querySelector('.section-header');
        var testBlocks = section.querySelectorAll('.test-block');
        var sectionMatches = false;

        testBlocks.forEach(function(tb) {
          var heading = tb.querySelector('.test-heading');
          var txt = heading ? heading.textContent.toLowerCase() : '';
          var match = !q || txt.indexOf(q) !== -1;
          tb.style.display = match ? '' : 'none';
          if (match) sectionMatches = true;
        });

        if (sectionHdr && !sectionMatches && q) {
          var secTxt = sectionHdr.textContent.toLowerCase();
          if (secTxt.indexOf(q) !== -1) {
            sectionMatches = true;
            testBlocks.forEach(function(tb) { tb.style.display = ''; });
          }
        }

        section.style.display = sectionMatches || !q ? '' : 'none';
      });

      updateCount();
    });

    updateCount();
  })();

  function clearFilter() {
    document.getElementById('uriFilter').value = '';
    document.getElementById('uriFilter').dispatchEvent(new Event('input'));
  }
"""


def report_name(time, ext):
    prefix = 'RedfishProtocolValidationReport'
    name = prefix + datetime.strftime(time, '_%m_%d_%Y_%H%M%S.' + ext)
    return name


def tsv_report(sut: SystemUnderTest, report_dir, time):
    file = report_dir / report_name(time, 'tsv')
    with open(str(file), 'w', encoding='utf-8') as fd:
        header = ('Assertion\tMethod\tStatus code\tURI\tResult\tMessage\t'
                  'Requirement\n')
        fd.write(header)
        for prefix, _ in sections:
            for assertion, results in sorted(
                    sut.results.items(), key=lambda x: x[0].name):
                if not assertion.name.startswith(prefix):
                    continue
                for r in results:
                    line = '{}\t{}\t{}\t{}\t{}\t{}\t{}\n'.format(
                        assertion.name, r['method'], r['status'], r['uri'],
                        r['result'].name, r['msg'], assertion.value)
                    fd.write(line)
    return str(file)


def _config_rows_html(args):
    config_rows_html = ""
    if args:
        for key, val in sorted(args.items()):
            if val is None:
                val = ""
            elif isinstance(val, list):
                val = " ".join(str(v) for v in val)
            else:
                val = str(val)
            if key == "password":
                val = "********" if val else ""
            config_rows_html += "<tr><td>{}</td><td>{}</td></tr>".format(
                html_mod.escape(key), html_mod.escape(val)
            )
    return config_rows_html


def html_report(sut: SystemUnderTest, report_dir, time, tool_version,
                 args=None):
    """
    Creates the HTML report for the system under test

    Args:
        sut: The system under test
        report_dir: The directory for the report
        time: The time the tests finished
        tool_version: The version of the tool
        args: A dict of the tool's CLI/config arguments, shown in the
              Configuration panel

    Returns:
        The path to the HTML report
    """
    file = report_dir / report_name(time, 'html')

    html = ''
    for prefix, section_name in sections:
        assertions = [a for a in sorted(sut.results.keys(),
                                        key=lambda x: x.name)
                      if a.name.startswith(prefix)]
        if not assertions:
            continue

        sec_pass = sec_warn = sec_fail = sec_skip = 0
        for assertion in assertions:
            for r in sut.results[assertion]:
                if r['result'] == Result.PASS:
                    sec_pass += 1
                elif r['result'] == Result.WARN:
                    sec_warn += 1
                elif r['result'] == Result.FAIL:
                    sec_fail += 1
                else:
                    sec_skip += 1
        sec_cnt = '<div class="section-counts">'
        if sec_pass:
            sec_cnt += '<span class="scnt scnt-pass">&#10003;&nbsp;{}</span>'.format(sec_pass)
        if sec_warn:
            sec_cnt += '<span class="scnt scnt-warn">&#9888;&nbsp;{}</span>'.format(sec_warn)
        if sec_fail:
            sec_cnt += '<span class="scnt scnt-fail">&#10007;&nbsp;{}</span>'.format(sec_fail)
        if sec_skip:
            sec_cnt += '<span class="scnt scnt-skip">&ndash;&nbsp;{}</span>'.format(sec_skip)
        sec_cnt += '</div>'

        html += '<div class="section">'
        html += '<div class="section-header"><span class="section-arrow">&#9660;</span>{}{}</div>'.format(
            html_mod.escape(section_name), sec_cnt
        )
        html += '<div class="section-body">'

        for assertion in assertions:
            results = sut.results[assertion]
            rvals = [r['result'] for r in results]
            if Result.FAIL in rvals:
                tb_cls = ' tb-fail'
            elif Result.WARN in rvals:
                tb_cls = ' tb-warn'
            elif Result.PASS in rvals:
                tb_cls = ' tb-pass'
            else:
                tb_cls = ''

            t_pass = sum(1 for r in results if r['result'] == Result.PASS)
            t_warn = sum(1 for r in results if r['result'] == Result.WARN)
            t_fail = sum(1 for r in results if r['result'] == Result.FAIL)
            t_skip = sum(1 for r in results if r['result'] == Result.NOT_TESTED)

            html += '<div class="test-block{}">'.format(tb_cls)
            html += '<div class="test-heading">'
            html += '<div class="test-heading-info">'
            html += '<div class="test-name">{}</div>'.format(html_mod.escape(assertion.name))
            html += '<div class="test-desc">{}</div>'.format(html_mod.escape(assertion.value))
            html += '<div class="result-chips">'
            if t_pass:
                html += '<span class="rchip rchip-pass">&#10003; {} Pass</span>'.format(t_pass)
            if t_warn:
                html += '<span class="rchip rchip-warn">&#9888; {} Warn</span>'.format(t_warn)
            if t_fail:
                html += '<span class="rchip rchip-fail">&#10007; {} Fail</span>'.format(t_fail)
            if t_skip:
                html += '<span class="rchip rchip-skip">&ndash; {} Not Tested</span>'.format(t_skip)
            html += '</div>'
            html += '</div>'  # test-heading-info
            html += '<span class="test-toggle">&#9660;</span>'
            html += '</div>'  # test-heading

            html += '<div class="test-body">'
            html += '<table class="test-results-table">'
            html += (
                '<thead><tr>'
                '<th class="col-method">Method</th>'
                '<th class="col-status">Status</th>'
                '<th class="col-uri">URI</th>'
                '<th class="col-result">Result</th>'
                '<th class="col-msg">Message</th>'
                '</tr></thead><tbody>'
            )
            for r in results:
                if r['result'] == Result.PASS:
                    badge = 'badge-pass'
                elif r['result'] == Result.WARN:
                    badge = 'badge-warn'
                elif r['result'] == Result.FAIL:
                    badge = 'badge-fail'
                else:
                    badge = 'badge-skip'
                result_label = r['result'].name.replace('_', ' ')
                html += (
                    '<tr><td>{}</td><td>{}</td><td>{}</td>'
                    '<td><span class="badge {}">{}</span></td><td>{}</td></tr>'
                ).format(
                    html_mod.escape(str(r['method'])),
                    html_mod.escape(str(r['status'])),
                    html_mod.escape(str(r['uri'])),
                    badge, result_label,
                    html_mod.escape(r['msg']),
                )
            html += '</tbody></table>'
            html += '</div>'  # test-body
            html += '</div>'  # test-block

        html += '</div>'  # section-body
        html += '</div>'  # section

    main_prefix = (
        '<div class="toolbar">'
        '<button class="toolbar-btn" onclick="expandAll()">&#9660;&nbsp; Expand All</button>'
        '<button class="toolbar-btn" onclick="collapseAll()">&#9654;&nbsp; Collapse All</button>'
        '</div>'
    )

    with open(str(file), 'w', encoding='utf-8') as fd:
        fd.write(
            build_html_report(
                page_title="Redfish Protocol Validator — Test Report",
                tool_title="Redfish Protocol Validator",
                filter_placeholder="Filter by assertion…",
                filter_count_label="test blocks",
                tool_link="https://github.com/DMTF/Redfish-Protocol-Validator",
                tool_repo="DMTF/Redfish-Protocol-Validator",
                tool_version=tool_version,
                generated_time=time.strftime('%c'),
                sut_host=html_mod.escape(str(sut.rhost)),
                sut_user=html_mod.escape(str(sut.username)),
                sut_password="********",
                sut_product=html_mod.escape(str(sut.product)),
                sut_manufacturer=html_mod.escape(str(sut.manufacturer)),
                sut_model=html_mod.escape(str(sut.model)),
                sut_firmware=html_mod.escape(str(sut.firmware_version)),
                pass_count=sut.summary_count(Result.PASS),
                warn_count=sut.summary_count(Result.WARN),
                fail_count=sut.summary_count(Result.FAIL),
                skip_count=sut.summary_count(Result.NOT_TESTED),
                sidebar_extra_html="",
                config_rows_html=_config_rows_html(args),
                extra_css=_RPV_EXTRA_CSS,
                main_prefix_html=main_prefix,
                main_content_html=html,
                extra_js=_RPV_EXTRA_JS,
            )
        )
    return str(file)


# ---------------------------------------------------------------------------
# XLSX helpers
# ---------------------------------------------------------------------------

def _thin_border():
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def _cell(ws, row, col, value="", bold=False, color=None, fill_hex=None,
          align="left", wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Segoe UI", size=10, bold=bold, color=color or "1A1A2E")
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = _thin_border()
    if fill_hex:
        c.fill = PatternFill("solid", fgColor=fill_hex)
    return c


def xlsx_report(sut: SystemUnderTest, report_dir, time, tool_version,
                 args=None):
    """
    Creates the Excel (xlsx) report for the system under test.

    Args:
        sut: The system under test
        report_dir: The directory for the report
        time: The time the tests finished
        tool_version: The version of the tool
        args: A dict of the tool's CLI/config arguments, shown in the
              Configuration section of the Summary sheet

    Returns:
        The path to the xlsx report
    """
    file = report_dir / report_name(time, 'xlsx')

    wb = openpyxl.Workbook()

    HDR_FILL = "0D1B2A"
    HDR_FONT = "FFFFFF"
    SUB_FILL = "1A3A5C"
    META_FILL = "F0F2F5"
    PASS_FILL = "D4EDDA"; PASS_FONT = "145A32"
    WARN_FILL = "FFF3CD"; WARN_FONT = "7D6008"
    FAIL_FILL = "F8D7DA"; FAIL_FONT = "7B241C"
    SKIP_FILL = "F5F7FA"; SKIP_FONT = "7F8C8D"
    SECTION_FILL = "1565C0"
    ASSERTION_FILL = "F7F9FC"
    SCORE_PASS = "27AE60"
    SCORE_WARN = "D68910"
    SCORE_FAIL = "C0392B"
    SCORE_SKIP = "5D6D7E"

    # ── Summary sheet ───────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"

    ws_sum.merge_cells("A1:D1")
    t = ws_sum["A1"]
    t.value = "Redfish Protocol Validator — Test Report"
    t.font = Font(name="Segoe UI", size=14, bold=True, color=HDR_FONT)
    t.fill = PatternFill("solid", fgColor=HDR_FILL)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 32

    ws_sum.merge_cells("A2:D2")
    s = ws_sum["A2"]
    s.value = "Version: {}    |    Generated: {}".format(tool_version, time.strftime("%c"))
    s.font = Font(name="Segoe UI", size=10, color=HDR_FONT)
    s.fill = PatternFill("solid", fgColor=SUB_FILL)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 18

    meta = [
        ("Target System", sut.rhost),
        ("User", sut.username),
        ("Product", sut.product),
        ("Manufacturer", sut.manufacturer),
        ("Model", sut.model),
        ("Firmware", sut.firmware_version),
    ]
    for i, (label, value) in enumerate(meta, start=3):
        _cell(ws_sum, i, 1, label, bold=True, color="6C757D", fill_hex=META_FILL, align="right")
        ws_sum.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        _cell(ws_sum, i, 2, value, fill_hex="FFFFFF")
        for col in range(2, 5):
            ws_sum.cell(row=i, column=col).border = _thin_border()

    r = 3 + len(meta) + 1
    score_labels = [("✓  PASS", SCORE_PASS), ("⚠  WARN", SCORE_WARN),
                    ("✗  FAIL", SCORE_FAIL), ("–  NOT TESTED", SCORE_SKIP)]
    for col, (label, fill) in enumerate(score_labels, start=1):
        _cell(ws_sum, r, col, label, bold=True, color=HDR_FONT,
              fill_hex=fill, align="center")
        ws_sum.row_dimensions[r].height = 22
    r += 1
    score_data = [
        (sut.summary_count(Result.PASS), PASS_FILL, PASS_FONT),
        (sut.summary_count(Result.WARN), WARN_FILL, WARN_FONT),
        (sut.summary_count(Result.FAIL), FAIL_FILL, FAIL_FONT),
        (sut.summary_count(Result.NOT_TESTED), SKIP_FILL, SKIP_FONT),
    ]
    for col, (count, fill, fnt) in enumerate(score_data, start=1):
        _cell(ws_sum, r, col, count, bold=True, color=fnt,
              fill_hex=fill, align="center")
        ws_sum.row_dimensions[r].height = 28
    r += 1

    if args:
        r += 1
        _cell(ws_sum, r, 1, "Configuration", bold=True, color=HDR_FONT,
              fill_hex=SUB_FILL, align="left")
        ws_sum.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        for key, val in sorted(args.items()):
            if val is None:
                val = ""
            elif isinstance(val, list):
                val = " ".join(str(v) for v in val)
            else:
                val = str(val)
            if key == "password":
                val = "********" if val else ""
            _cell(ws_sum, r, 1, key, bold=True, color="6C757D", fill_hex=META_FILL, align="right")
            ws_sum.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            _cell(ws_sum, r, 2, val, fill_hex="FFFFFF", wrap=True)
            for col in range(2, 5):
                ws_sum.cell(row=r, column=col).border = _thin_border()
            r += 1

    ws_sum.column_dimensions["A"].width = 20
    for col_letter in ["B", "C", "D"]:
        ws_sum.column_dimensions[col_letter].width = 24
    ws_sum.sheet_view.showGridLines = False

    # ── Results sheet ───────────────────────────────────────────────────
    ws = wb.create_sheet("Results")

    col_headers = ["Section", "Assertion", "Requirement", "Method",
                  "Status Code", "URI", "Result", "Message"]
    col_widths = [18, 22, 40, 10, 12, 40, 12, 55]

    for col, (hdr, width) in enumerate(zip(col_headers, col_widths), start=1):
        _cell(ws, 1, col, hdr, bold=True, color=HDR_FONT,
              fill_hex=SUB_FILL, align="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    row = 2
    result_style = {
        Result.PASS: (PASS_FILL, PASS_FONT),
        Result.WARN: (WARN_FILL, WARN_FONT),
        Result.FAIL: (FAIL_FILL, FAIL_FONT),
    }

    for prefix, section_name in sections:
        assertions = [a for a in sorted(sut.results.keys(), key=lambda x: x.name)
                      if a.name.startswith(prefix)]
        for assertion in assertions:
            for r in sut.results[assertion]:
                fill, fnt = result_style.get(r['result'], (SKIP_FILL, SKIP_FONT))
                result_label = r['result'].name.replace('_', ' ')

                _cell(ws, row, 1, section_name, fill_hex=SECTION_FILL, color=HDR_FONT, wrap=True)
                _cell(ws, row, 2, assertion.name, fill_hex=ASSERTION_FILL, bold=True, wrap=True)
                _cell(ws, row, 3, assertion.value, fill_hex=ASSERTION_FILL, wrap=True)
                _cell(ws, row, 4, r['method'], wrap=True)
                _cell(ws, row, 5, r['status'], align="center")
                _cell(ws, row, 6, r['uri'], wrap=True)
                _cell(ws, row, 7, result_label, fill_hex=fill, color=fnt,
                      bold=True, align="center")
                _cell(ws, row, 8, r['msg'], wrap=True)
                ws.row_dimensions[row].height = 30
                row += 1

    if row > 2:
        ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(col_headers)), row - 1)

    wb.save(str(file))
    return str(file)


def json_results(sut: SystemUnderTest, report_dir, time, tool_version):
    file = report_dir / 'results.json'
    results = {
        'ToolName': 'Redfish-Protocol-Validator v%s' % tool_version,
        'Timestamp': {
            'DateTime': '{:%Y-%m-%dT%H:%M:%S%Z}'.format(time)
        },
        'Service': {
            'BaseURL': sut.rhost,
            'Manufacturer': sut.manufacturer,
            'Product': sut.product,
            'Model': sut.model,
            'FirmwareVersion': sut.firmware_version
        },
        'TestResults': {
            'Protocol Validations': {
                'pass': sut.summary_count(Result.PASS),
                'fail': sut.summary_count(Result.FAIL),
                'skip': sut.summary_count(Result.NOT_TESTED),
                'warn': sut.summary_count(Result.WARN)
            },
            'ErrorMessages': []
        }
    }
    with open(str(file), 'w', encoding='utf-8') as fd:
        json.dump(results, fd, indent=4)
